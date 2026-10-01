"""Alunas do Club que sao candidatas a subir para o PRO.

Criterio, na ordem em que e aplicado:

1. Produto = Moonshot Club, e nao ha registro dela em PRO ou Elite.
2. Data de entrada a partir de 08/2026 - regra do Marcelo: entrada no Club
   so passou a existir nesse mes.
3. Corte do Marcelo: fora quem fatura abaixo de 15k.

Faturamento vazio nao e faturamento abaixo de 15k - e faturamento nao medido,
e nao da para cortar por um numero que ninguem escreveu. Por isso a planilha
tem duas abas:

  'Club -> Pro'      quem tem faturamento registrado de 15k para cima
  'Sem faturamento'  quem nao tem faturamento em nenhuma das duas planilhas

A segunda aba nao foi cortada porque nao ha como saber se ela cumpre ou nao
o corte. Preencher o faturamento dessas alunas decide para qual lado elas vao.
"""
import re
import unicodedata
import warnings

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')

CONSULTORIAS = 'dados/controle_consultorias_live.xlsx'
MATRICULAS = 'dados/matriculados_club_live.xlsx'
SAIDA = 'entregaveis/ascensao_club_para_pro.xlsx'
CORTE_CLUB = pd.Timestamp('2026-08-01')
CORTE_FAT = 15000


def chave(n):
    n = unicodedata.normalize('NFKD', str(n)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z ]', ' ', n.lower())).strip()


# A grafia da aba do Controle Consultorias e a oficial: a aba e preenchida pelo
# proprio consultor, o campo da matricula e digitado por terceiro. Sem isso
# 'Anameli' e 'Anameli' viram dois consultores na mesma lista.
CANONICO = {}
for _aba in openpyxl.load_workbook(CONSULTORIAS, read_only=True).sheetnames:
    _n = re.sub(r'\s+Geral\s*$', '', _aba.strip())
    if chave(_n):
        CANONICO.setdefault(chave(_n), _n)


def telefone(v):
    """O telefone vem de celulas que o Sheets tratou como numero, entao chega
    como '19999594808.0'. Formata so o que tem cara de numero brasileiro:
    11 digitos com 9 na frente do numero, ou 10 digitos de fixo. O resto
    (numero de fora, numero truncado) sai como esta na origem - por DDD de
    palpite nao se liga para ninguem."""
    if v is None or pd.isna(v):
        return ''
    t = str(v).strip()
    if not t:
        return ''
    if '/' in t:  # algumas celulas guardam dois numeros separados por barra
        return ' / '.join(x for x in (telefone(p) for p in t.split('/')) if x)
    t = re.sub(r'\.0$', '', t)
    d = re.sub(r'\D', '', t)
    if len(d) in (12, 13) and d.startswith('55'):
        d = d[2:]
    if len(d) == 11 and d[2] == '9':
        return f'({d[:2]}) {d[2:7]}-{d[7:]}'
    if len(d) == 10:
        return f'({d[:2]}) {d[2:6]}-{d[6:]}'
    return t


def consultor(v):
    # 'Thiago Geral' e o nome da ABA, nao do consultor.
    if v is None or pd.isna(v) or not str(v).strip():
        return '(sem consultor)'
    n = re.sub(r'\s+Geral\s*$', '', str(v).strip())
    return CANONICO.get(chave(n), n)


def observacoes():
    """Duas colunas de texto livre descrevem a aluna: 'SITUACAO ALUNA' nas abas
    dos consultores e 'Observacoes' nas abas de matricula. Procura as duas pelo
    nome do cabecalho, nunca por posicao - as abas tem de 7 a 18 colunas e a
    posicao fixa cai em cima de data de consultoria."""
    m = {}
    for arq, chave_nome, chave_obs in (
            (CONSULTORIAS, ('aluna',), 'situacao'),
            (MATRICULAS, ('coluna', 'nome', 'aluna', 'aluno', 'alunas'), 'observac')):
        for aba in openpyxl.load_workbook(arq, read_only=True).sheetnames:
            cn = co = None
            for cab in (0, 1, 2):
                d = pd.read_excel(arq, sheet_name=aba, header=cab)
                cols = {chave(c): c for c in d.columns}
                cn = next((v for k, v in cols.items()
                           if any(k.startswith(p) for p in chave_nome)), None)
                co = next((v for k, v in cols.items() if k.startswith(chave_obs)), None)
                if cn and co:
                    break
            if not (cn and co):
                continue
            for _, r in d.iterrows():
                nm, o = chave(r[cn]), r[co]
                if not nm or 'moonshot' in nm or pd.isna(o):
                    continue
                o = str(o).strip()
                # data solta nessa coluna e erro de preenchimento, nao observacao
                if not o or o.lower() in ('nan', '-') or re.fullmatch(r'\d{4}-\d{2}-\d{2}.*', o):
                    continue
                if nm not in m or len(o) > len(m[nm]):
                    m[nm] = o
    return m


a = pd.read_pickle('/tmp/claude-0/-home-user-moonshot-analise/'
                   'aa4a53bc-6d8a-5522-8729-cf6ad6d2e048/scratchpad/asc.pkl')
entrada = pd.to_datetime(a.ini_mat.fillna(a.ini_cons), errors='coerce')
c = a[a.e_club & ~a.ja_pro & (entrada >= CORTE_CLUB)].copy()
c['entrada'] = entrada[c.index]

OBS = observacoes()
S = pd.DataFrame({
    'Aluna': c.nome.str.strip(),
    'Contato': c.tel.map(telefone),
    'Consultor Responsável': c.cons_aba.fillna(c.cons_mat).map(consultor),
    'Faturamento': c.fat,
    'Data de Entrada': c.entrada.dt.strftime('%d/%m/%Y'),
    'Cidade': c.cidade.fillna('').str.strip(),
    'Ramo': c.ramo.fillna('').str.strip(),
    'Observações sobre a aluna': c.nome.map(lambda n: OBS.get(chave(n), '')),
})
# Faturamento desconhecido vai para o fim, nao para o meio como se fosse zero.
S = S.sort_values('Faturamento', ascending=False, na_position='last').reset_index(drop=True)

ACIMA = S[S.Faturamento >= CORTE_FAT].reset_index(drop=True)
SEM = S[S.Faturamento.isna()].drop(columns='Faturamento').reset_index(drop=True)

with pd.ExcelWriter(SAIDA, engine='xlsxwriter') as xw:
    wb = xw.book
    cab = wb.add_format({'bold': True, 'bg_color': '#1F2937', 'font_color': 'white',
                         'border': 1, 'align': 'center', 'valign': 'vcenter',
                         'text_wrap': True})
    txt = wb.add_format({'valign': 'top', 'border': 1})
    wrap = wb.add_format({'valign': 'top', 'border': 1, 'text_wrap': True})
    din = wb.add_format({'valign': 'top', 'border': 1, 'num_format': 'R$ #,##0'})
    vazio = wb.add_format({'valign': 'top', 'border': 1, 'font_color': '#9CA3AF',
                           'align': 'center'})
    larg = {'Aluna': 38, 'Contato': 20, 'Consultor Responsável': 18, 'Faturamento': 15,
            'Data de Entrada': 15, 'Cidade': 22, 'Ramo': 24,
            'Observações sobre a aluna': 55}

    for aba, D in (('Club → Pró', ACIMA), ('Sem faturamento', SEM)):
        ws = wb.add_worksheet(aba)
        for j, nome in enumerate(D.columns):
            ws.write(0, j, nome, cab)
            ws.set_column(j, j, larg[nome])
        for i, row in D.iterrows():
            for j, nome in enumerate(D.columns):
                v = row[nome]
                if nome == 'Faturamento':
                    ws.write_number(i + 1, j, float(v), din)
                elif v == '' or pd.isna(v):
                    ws.write(i + 1, j, '—', vazio)
                else:
                    ws.write(i + 1, j, str(v),
                             wrap if nome == 'Observações sobre a aluna' else txt)
        ws.set_row(0, 32)
        ws.freeze_panes(1, 1)
        ws.autofilter(0, 0, max(len(D), 1), len(D.columns) - 1)

print('Club entrando a partir de 08/2026:', len(S))
print('  >= 15k (aba Club -> Pro):', len(ACIMA))
print('  < 15k (cortadas):', int((S.Faturamento < CORTE_FAT).sum()))
print('  sem faturamento (aba Sem faturamento):', len(SEM))
for aba, D in (('Club -> Pro', ACIMA), ('Sem faturamento', SEM)):
    print(f'\n[{aba}] {len(D)} alunas'
          f' | telefone {int((D.Contato != "").sum())}'
          f' | cidade {int((D.Cidade != "").sum())}'
          f' | ramo {int((D.Ramo != "").sum())}'
          f' | observacao {int((D["Observações sobre a aluna"] != "").sum())}')
    print(D.groupby('Consultor Responsável').size().sort_values(ascending=False).to_string())
